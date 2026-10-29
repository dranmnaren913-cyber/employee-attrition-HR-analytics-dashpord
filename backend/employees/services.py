import csv
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse
from django.db import transaction
from .models import Employee, ImportLog
from departments.models import Department, JobRole

def parse_date(date_str):
    if not date_str:
        return datetime.today().date()
    if isinstance(date_str, datetime):
        return date_str.date()
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%m/%d/%Y', '%d/%m/%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(str(date_str).strip(), fmt).date()
        except ValueError:
            pass
    return datetime.today().date()

def process_employee_import(file_obj, user=None):
    filename = file_obj.name.lower()
    records = []
    
    # 1. Read records from CSV or Excel
    if filename.endswith('.csv'):
        try:
            decoded_file = file_obj.read().decode('utf-8-sig', errors='replace')
            io_string = io.StringIO(decoded_file)
            reader = csv.DictReader(io_string)
            for row in reader:
                records.append({k.strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k})
        except Exception as e:
            return {
                'success': False,
                'error': f"Failed to read CSV file: {str(e)}",
                'total_records': 0,
                'successful_records': 0,
                'failed_records': 0,
                'errors': [{'row': 0, 'error': f"File read error: {str(e)}"}]
            }
    elif filename.endswith(('.xlsx', '.xls')):
        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            sheet = wb.active
            headers = [cell.value for cell in sheet[1]]
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if any(row):
                    row_dict = {}
                    for h, v in zip(headers, row):
                        if h:
                            row_dict[str(h).strip()] = v
                    records.append(row_dict)
        except Exception as e:
            return {
                'success': False,
                'error': f"Failed to read Excel file: {str(e)}",
                'total_records': 0,
                'successful_records': 0,
                'failed_records': 0,
                'errors': [{'row': 0, 'error': f"File read error: {str(e)}"}]
            }
    else:
        return {
            'success': False,
            'error': 'Unsupported file format. Please upload a CSV or Excel (.xlsx) file.',
            'total_records': 0,
            'successful_records': 0,
            'failed_records': 0,
            'errors': []
        }

    total_records = len(records)
    successful_records = 0
    failed_records = 0
    errors = []

    # Map of cached departments and job roles
    dept_cache = {d.department_name.lower(): d for d in Department.objects.all()}
    role_cache = {(r.role_name.lower(), r.department_id): r for r in JobRole.objects.all()}

    # Check existing IDs and emails
    existing_emp_ids = set(Employee.objects.values_list('employee_id', flat=True))
    existing_emails = set(Employee.objects.values_list('email', flat=True))

    valid_employees_to_create = []

    for idx, row in enumerate(records, start=2):
        row_errors = []
        
        # Normalize key names
        row_norm = {k.lower().replace(' ', '_').replace('-', '_'): v for k, v in row.items()}
        
        emp_id = str(row_norm.get('employee_id') or row_norm.get('id') or '').strip()
        name = str(row_norm.get('name') or row_norm.get('full_name') or '').strip()
        email = str(row_norm.get('email') or '').strip()

        if not emp_id:
            row_errors.append("Missing Employee ID")
        elif emp_id in existing_emp_ids:
            row_errors.append(f"Duplicate Employee ID: {emp_id}")

        if not name:
            row_errors.append("Missing Name")

        if not email:
            row_errors.append("Missing Email")
        elif email in existing_emails:
            row_errors.append(f"Duplicate Email: {email}")

        # Department resolution
        dept_name = str(row_norm.get('department') or row_norm.get('department_name') or 'General').strip()
        dept_key = dept_name.lower()
        if dept_key not in dept_cache:
            dept = Department.objects.create(
                department_name=dept_name,
                code=dept_name[:3].upper() + str(Department.objects.count() + 1)
            )
            dept_cache[dept_key] = dept
        dept = dept_cache[dept_key]

        # Job role resolution
        role_name = str(row_norm.get('job_role') or row_norm.get('role') or row_norm.get('role_name') or 'Specialist').strip()
        role_key = (role_name.lower(), dept.id)
        if role_key not in role_cache:
            role = JobRole.objects.create(
                role_name=role_name,
                department=dept
            )
            role_cache[role_key] = role
        role = role_cache[role_key]

        # Numerical fields with validation
        try:
            age = int(float(row_norm.get('age', 30)))
            if age < 18 or age > 80:
                row_errors.append(f"Invalid age ({age}), must be between 18 and 80")
        except (ValueError, TypeError):
            row_errors.append("Age must be a valid number")
            age = 30

        try:
            experience = float(row_norm.get('experience', 2.0))
            if experience < 0:
                row_errors.append("Experience cannot be negative")
        except (ValueError, TypeError):
            row_errors.append("Experience must be a valid number")
            experience = 2.0

        try:
            salary = float(row_norm.get('salary', 50000.0))
            if salary <= 0:
                row_errors.append("Salary must be positive")
        except (ValueError, TypeError):
            row_errors.append("Salary must be a valid number")
            salary = 50000.0

        try:
            perf_score = int(float(row_norm.get('performance_score', 3)))
            perf_score = max(1, min(5, perf_score))
        except (ValueError, TypeError):
            perf_score = 3

        try:
            sat_score = int(float(row_norm.get('satisfaction_score', 3)))
            sat_score = max(1, min(5, sat_score))
        except (ValueError, TypeError):
            sat_score = 3

        try:
            wlb = int(float(row_norm.get('work_life_balance', 3)))
            wlb = max(1, min(5, wlb))
        except (ValueError, TypeError):
            wlb = 3

        gender = str(row_norm.get('gender') or 'Male').capitalize()
        if gender not in ['Male', 'Female', 'Other']:
            gender = 'Other'

        education = str(row_norm.get('education') or 'Bachelor')
        phone = str(row_norm.get('phone') or '')
        joining_date = parse_date(row_norm.get('joining_date'))

        overtime = str(row_norm.get('overtime') or 'No').capitalize()
        overtime = 'Yes' if overtime in ('Yes', 'Y', 'True', '1') else 'No'

        emp_status = str(row_norm.get('employment_status') or 'Active').capitalize()
        if emp_status not in ['Active', 'On Leave', 'Resigned', 'Terminated']:
            emp_status = 'Active'

        attrition = str(row_norm.get('attrition_status') or 'No').capitalize()
        attrition = 'Yes' if attrition in ('Yes', 'Y', 'True', '1') or emp_status in ('Resigned', 'Terminated') else 'No'

        if row_errors:
            failed_records += 1
            errors.append({
                'row': idx,
                'employee_id': emp_id or f"Row {idx}",
                'errors': row_errors
            })
        else:
            valid_employees_to_create.append(
                Employee(
                    employee_id=emp_id,
                    name=name,
                    email=email,
                    phone=phone,
                    gender=gender,
                    age=age,
                    education=education,
                    department=dept,
                    job_role=role,
                    joining_date=joining_date,
                    experience=experience,
                    salary=salary,
                    performance_score=perf_score,
                    satisfaction_score=sat_score,
                    work_life_balance=wlb,
                    overtime=overtime,
                    employment_status=emp_status,
                    attrition_status=attrition
                )
            )
            existing_emp_ids.add(emp_id)
            existing_emails.add(email)

    # Bulk insert valid employees inside a transaction
    if valid_employees_to_create:
        with transaction.atomic():
            Employee.objects.bulk_create(valid_employees_to_create)
            successful_records = len(valid_employees_to_create)

    # Log to ImportLog
    import_log = ImportLog.objects.create(
        filename=getattr(file_obj, 'name', 'import_file'),
        total_records=total_records,
        successful_records=successful_records,
        failed_records=failed_records,
        uploaded_by=user if user and user.is_authenticated else None,
        error_details=errors[:50] # Top 50 errors
    )

    return {
        'success': True,
        'import_id': import_log.id,
        'total_records': total_records,
        'successful_records': successful_records,
        'failed_records': failed_records,
        'errors': errors
    }


def export_employees(queryset, export_format='csv'):
    """
    Exports employee queryset to CSV or Excel
    """
    headers = [
        'Employee ID', 'Name', 'Email', 'Phone', 'Gender', 'Age',
        'Department', 'Job Role', 'Education', 'Joining Date',
        'Experience (Years)', 'Salary ($)', 'Performance Score (1-5)',
        'Satisfaction Score (1-5)', 'Work-Life Balance (1-5)', 'Overtime',
        'Employment Status', 'Attrition Status', 'Risk Level', 'Attrition Probability (%)'
    ]

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

    if export_format.lower() == 'excel' or export_format.lower() == 'xlsx':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Employees"

        # Styling
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
        border_thin = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for emp in queryset.select_related('department', 'job_role').prefetch_related('predictions'):
            pred = emp.predictions.order_by('-prediction_date').first()
            risk = pred.risk_level if pred else ('High' if emp.satisfaction_score <= 2 else 'Low')
            prob = f"{pred.probability:.1f}%" if pred else ('75.0%' if emp.satisfaction_score <= 2 else '18.0%')

            row_data = [
                emp.employee_id,
                emp.name,
                emp.email,
                emp.phone or 'N/A',
                emp.gender,
                emp.age,
                emp.department.department_name,
                emp.job_role.role_name,
                emp.education,
                emp.joining_date.strftime('%Y-%m-%d') if emp.joining_date else '',
                emp.experience,
                float(emp.salary),
                emp.performance_score,
                emp.satisfaction_score,
                emp.work_life_balance,
                emp.overtime,
                emp.employment_status,
                emp.attrition_status,
                risk,
                prob
            ]
            ws.append(row_data)

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="employees_export_{timestamp}.xlsx"'
        return response

    else: # Default CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="employees_export_{timestamp}.csv"'

        writer = csv.writer(response)
        writer.writerow(headers)

        for emp in queryset.select_related('department', 'job_role').prefetch_related('predictions'):
            pred = emp.predictions.order_by('-prediction_date').first()
            risk = pred.risk_level if pred else ('High' if emp.satisfaction_score <= 2 else 'Low')
            prob = f"{pred.probability:.1f}%" if pred else ('75.0%' if emp.satisfaction_score <= 2 else '18.0%')

            writer.writerow([
                emp.employee_id,
                emp.name,
                emp.email,
                emp.phone or 'N/A',
                emp.gender,
                emp.age,
                emp.department.department_name,
                emp.job_role.role_name,
                emp.education,
                emp.joining_date.strftime('%Y-%m-%d') if emp.joining_date else '',
                emp.experience,
                float(emp.salary),
                emp.performance_score,
                emp.satisfaction_score,
                emp.work_life_balance,
                emp.overtime,
                emp.employment_status,
                emp.attrition_status,
                risk,
                prob
            ])

        return response
